#!/usr/bin/env python3
"""Paper A manuscript — offline build without pandoc.

manuscript.md (pandoc-style [@key] citations) + references.bib
  -> export/manuscript_rendered.md   author–year citations and a reference list
  -> export/manuscript.docx          python-docx (headings, lists, tables, figures)
  -> export/manuscript.pdf           LibreOffice headless conversion of the .docx (if installed)

Every cited key must exist in references.bib (the build fails otherwise).

    python docs/research/paper/manuscript/tools/build.py
"""
import os
import re
import shutil
import subprocess
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
MS_DIR = os.path.abspath(os.path.join(HERE, ".."))
BUILD = os.path.join(MS_DIR, "export")
SRC = os.path.join(MS_DIR, "manuscript.md")
BIB = os.path.join(MS_DIR, "references.bib")

ACCENTS = [(r"{\`e}", "è"), (r"{\'e}", "é"), (r"{\'o}", "ó"), (r"{\v{Z}}", "Ž"), (r"{\c{S}}", "Ş"), (r"{\c{s}}", "ş"),
           (r'{\"a}', "ä"), (r"\'e", "é"), ("--", "–")]


def delatex(s):
    for a, b in ACCENTS:
        s = s.replace(a, b)
    s = s.replace("\\_", "_")
    return re.sub(r"[{}]", "", s).strip()


def parse_bib(path):
    txt = open(path, encoding="utf-8").read()
    entries = {}
    for m in re.finditer(r"@(\w+)\{([^,]+),(.*?)\n\}", txt, re.S):
        typ, key, body = m.group(1).lower(), m.group(2).strip(), m.group(3)
        fields = {}
        for fm in re.finditer(r"(\w+)\s*=\s*\{(.*?)\}\s*,?\s*$", body, re.S | re.M):
            fields[fm.group(1).lower()] = fm.group(2)
        fields["_type"] = typ
        entries[key] = fields
    return entries


def surnames(author_field):
    if not author_field:
        return []
    out = []
    for a in re.split(r"\s+and\s+", author_field):
        a = a.strip()
        if a == "others":
            out.append("others")
        elif "," in a:
            out.append(delatex(a.split(",")[0]))
        else:
            out.append(delatex(a.split()[-1] if not a.startswith("{") else a))
    return out


def cite_label(e):
    sn = surnames(e.get("author", ""))
    if not sn:
        return delatex(e.get("title", "?")), e.get("year", "")
    if len(sn) == 1:
        lab = sn[0]
    elif len(sn) == 2 and sn[1] != "others":
        lab = f"{sn[0]} and {sn[1]}"
    else:
        lab = f"{sn[0]} et al."
    return lab, e.get("year", "")


def fmt_authors(field):
    if not field:
        return ""
    parts = []
    for a in re.split(r"\s+and\s+", field):
        a = a.strip()
        if a == "others":
            parts.append("et al.")
            continue
        if "," in a:
            last, first = a.split(",", 1)
            ini = " ".join(w[0] + "." for w in re.split(r"[\s-]+", delatex(first)) if w and w[0].isalpha())
            parts.append(f"{delatex(last)} {ini}".strip())
        else:
            parts.append(delatex(a))
    if parts and parts[-1] == "et al.":
        return ", ".join(parts[:-1]) + ", et al."
    return ", ".join(parts)


def fmt_ref(key, e):
    a = fmt_authors(e.get("author", ""))
    title = delatex(e.get("title", ""))
    yr = e.get("year", "")
    venue = delatex(e.get("journal", "") or e.get("booktitle", "") or e.get("howpublished", "") or e.get("publisher", ""))
    vol = e.get("volume", "")
    num = e.get("number", "")
    pages = delatex(e.get("pages", ""))
    s = f"{a} ({yr}). {title}." if a else f"{title} ({yr})."
    if venue:
        s += f" *{venue}*"
        if vol:
            s += f" {vol}"
            if num:
                s += f"({delatex(num)})"
        if pages:
            s += f": {pages}"
        s += "."
    if e.get("doi"):
        s += f" doi:{e['doi']}."
    if e.get("eprint"):
        s += f" arXiv:{e['eprint']}."
    if e.get("note") and "Given names" not in e.get("note", ""):
        s += f" {delatex(e['note'])}."
    return s


def render_citations(md, bib):
    used = []

    def lab(key):
        if key not in bib:
            raise SystemExit(f"citation key not in references.bib: {key}")
        if key not in used:
            used.append(key)
        return cite_label(bib[key])

    def bracket(m):
        keys = [k.strip().lstrip("@") for k in m.group(1).split(";")]
        return "(" + "; ".join(f"{a}, {y}" for a, y in (lab(k) for k in keys)) + ")"
    md = re.sub(r"\[(@[\w\-]+(?:\s*;\s*@[\w\-]+)*)\]", bracket, md)

    def narrative(m):
        a, y = lab(m.group(1))
        return f"{a} ({y})"
    md = re.sub(r"(?<![\w\[/])@([a-z]+\d{4}[a-z0-9]*)", narrative, md)
    return md, used


def build_md():
    md = open(SRC, encoding="utf-8").read()
    bib = parse_bib(BIB)
    fm, body = md.split("---", 2)[1], md.split("---", 2)[2]
    meta = {}
    cur = None
    for ln in fm.splitlines():
        m = re.match(r"^(\S[^:]*):\s*(.*)$", ln)
        if m:
            cur = m.group(1)
            meta[cur] = m.group(2).strip().strip('"')
            if not meta[cur]:
                meta[cur] = []
        elif ln.strip().startswith("- ") and cur:
            meta[cur].append(ln.strip()[2:].strip('"'))
    body = body.replace("<!-- Rendered by tools/build.py from references.bib (pandoc-style [@key] citations). -->", "@@REFLIST@@")
    body, used = render_citations(body, bib)
    refs = sorted(used, key=lambda k: (cite_label(bib[k])[0].lower(), bib[k].get("year", "")))
    reflist = "\n".join(f"- {fmt_ref(k, bib[k])}" for k in refs)
    body = body.replace("@@REFLIST@@", reflist)
    body = re.sub(r"<!-- (BEGIN|END) TABLE \w+ -->\n", "", body)
    head = f"# {meta['title']}\n\n" + "\n\n".join(f"*{a}*" for a in meta["author"]) + \
           f"\n\n*Target venue: {meta['target-venue']}*\n\n*{meta['status']}*\n"
    out = head + body
    os.makedirs(BUILD, exist_ok=True)
    with open(os.path.join(BUILD, "manuscript_rendered.md"), "w", encoding="utf-8", newline="\n") as f:
        f.write(out)
    unused = sorted(set(bib) - set(used))
    return out, used, unused


# ------------------------------------------------------------------ docx
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)")


def add_inline(par, text, size=None, bold=False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = par.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = par.add_run(part[1:-1])
            run.font.name = "Consolas"
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = par.add_run(part[1:-1])
            run.italic = True
            run.bold = bold
        else:
            run = par.add_run(part)
            run.bold = bold
        if size:
            run.font.size = Pt(size)


def add_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows if not re.match(r"^\|[\s\-:|]+\|$", r.strip())]
    ncol = max(len(r) for r in cells)
    t = doc.add_table(rows=len(cells), cols=ncol)
    t.style = "Table Grid"
    size = 6.5 if ncol > 5 else 7.5
    for i, r in enumerate(cells):
        for j in range(ncol):
            par = t.cell(i, j).paragraphs[0]
            add_inline(par, r[j] if j < len(r) else "", size=size, bold=(i == 0))
    doc.add_paragraph()


def build_docx(md):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.9)
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", ln)
        if m:
            h = doc.add_heading(level=min(len(m.group(1)), 3) if len(m.group(1)) > 1 else 0 if i == 0 else 1)
            add_inline(h, m.group(2))
            i += 1
            continue
        m = re.match(r"^!\[(.*)\]\((.*)\)$", ln.strip())
        if m:
            img = os.path.join(MS_DIR, m.group(2))
            if os.path.exists(img):
                doc.add_picture(img, width=Inches(6.4))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap = doc.add_paragraph()
            add_inline(cap, m.group(1), size=9)
            i += 1
            continue
        if ln.strip() == "---":
            i += 1
            continue
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", ln)
        if m:
            text = m.group(3)
            i += 1
            while i < len(lines) and lines[i].startswith("  ") and lines[i].strip() and not re.match(r"^\s*([-*]|\d+\.)\s", lines[i]):
                text += " " + lines[i].strip()
                i += 1
            if m.group(2) in "-*":
                p = doc.add_paragraph(style="List Bullet 2" if len(m.group(1)) >= 2 else "List Bullet")
                add_inline(p, text)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.first_line_indent = Inches(-0.25)
                add_inline(p, f"{m.group(2)} " + text)
            continue
        if not ln.strip():
            i += 1
            continue
        text = ln.strip()
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#|\||!\[|\s*([-*]|\d+\.)\s|---)", lines[i]):
            text += " " + lines[i].strip()
            i += 1
        p = doc.add_paragraph()
        add_inline(p, text.replace("  ", " "))
    path = os.path.join(BUILD, "manuscript.docx")
    doc.save(path)
    return path


def build_pdf(docx_path):
    cands = [shutil.which("soffice"), r"C:\Program Files\LibreOffice\program\soffice.exe"]
    exe = next((c for c in cands if c and os.path.exists(c)), None)
    if not exe:
        print("LibreOffice not found: PDF skipped")
        return None
    subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", BUILD, docx_path], check=True,
                   capture_output=True, timeout=300)
    return os.path.join(BUILD, "manuscript.pdf")


def main():
    md, used, unused = build_md()
    print(f"citations: {len(used)} keys cited; unused bib entries: {unused}")
    d = build_docx(md)
    print("wrote", d)
    p = build_pdf(d)
    if p:
        print("wrote", p)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
